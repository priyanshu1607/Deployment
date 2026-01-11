from docx import Document
import os

# Read the existing DOCX to understand its structure
existing_doc = Document("AI Deployment.docx")

print("=" * 80)
print("EXISTING DOCX STRUCTURE")
print("=" * 80)

for i, para in enumerate(existing_doc.paragraphs[:20]):  # First 20 paragraphs
    if para.text.strip():
        print(f"\nParagraph {i}:")
        print(f"  Style: {para.style.name}")
        print(f"  Text: {para.text[:100]}")  # First 100 chars

print("\n" + "=" * 80)
print("TABLES IN DOCUMENT")
print("=" * 80)
print(f"Number of tables: {len(existing_doc.tables)}")

for i, table in enumerate(existing_doc.tables[:3]):  # First 3 tables
    print(f"\nTable {i}: {len(table.rows)} rows x {len(table.columns)} columns")
    if table.rows:
        print(f"  First row: {[cell.text[:30] for cell in table.rows[0].cells]}")
